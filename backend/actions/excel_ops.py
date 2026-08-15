"""
Excel operations via openpyxl — read cells and write values/formulas.
"""

from __future__ import annotations

import logging

log = logging.getLogger("addled.excel")


class ExcelOps:
    """Read/write .xlsx workbooks (keeps formulas and formatting intact)."""

    async def read(self, path: str, sheet: str | None = None,
                   max_rows: int = 500) -> dict:
        """Return rows as a list of lists (values only)."""
        if not path:
            return {"success": False, "error": "path is required"}
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            ws = wb[sheet] if sheet else wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if v is None else v for v in row])
                if len(rows) >= max_rows:
                    break
            wb.close()
            return {"success": True, "sheet": ws.title,
                    "rows": rows, "rowCount": len(rows)}
        except ImportError:
            return {"success": False,
                    "error": "openpyxl not installed. Run: pip install openpyxl"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    async def write(self, path: str, sheet: str = "Sheet1",
                    cells: list | None = None, grid: list | None = None) -> dict:
        """
        Write cell values. `cells` = [{row, col, value}...]; `grid` = matrix
        written starting at A1. Creates the file if it doesn't exist.
        """
        if not path:
            return {"success": False, "error": "path is required"}
        try:
            import openpyxl
            import os
            if os.path.isfile(path):
                wb = openpyxl.load_workbook(path)
            else:
                wb = openpyxl.Workbook()
            ws = wb[sheet] if sheet in wb.sheetnames else wb.active
            if ws.title != sheet:
                ws.title = sheet

            written = 0
            if grid:
                for r, row in enumerate(grid, start=1):
                    for c, value in enumerate(row, start=1):
                        ws.cell(row=r, column=c, value=value)
                        written += 1
            for cell in cells or []:
                r = cell.get("row")
                c = cell.get("col") or cell.get("column")
                if r is None or c is None:
                    continue
                ws.cell(row=int(r), column=int(c), value=cell.get("value"))
                written += 1

            wb.save(path)
            return {"success": True, "path": path, "sheet": ws.title,
                    "cellsWritten": written}
        except ImportError:
            return {"success": False,
                    "error": "openpyxl not installed. Run: pip install openpyxl"}
        except Exception as e:
            return {"success": False, "error": str(e)}


excel_ops = ExcelOps()
